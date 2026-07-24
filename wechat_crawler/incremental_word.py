from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NOISE_LINES = {
    "收藏",
    "听过",
    "知道了",
    "喜欢此内容的人还喜欢",
    "微信扫一扫",
    "使用小程序",
    "使用完整服务",
    "取消",
    "允许",
    "分析",
    "已同步到看一看",
    "赞",
    "在看",
    "分享",
    "留言",
}
CJK_FONT = "Arial Unicode MS"


class IncrementalWordExporter:
    def export_account(
        self,
        articles: Sequence[Mapping[str, Any]],
        output_path: str | Path,
        *,
        account_name: str,
        mode: str,
        cycle_key: str,
        generated_at: datetime | None = None,
    ) -> Path:
        if not articles:
            raise ValueError("cannot export an empty account document")
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(target.suffix + ".tmp")
        generated = generated_at or datetime.now().astimezone()
        period_label = self._article_period_label(articles, cycle_key)

        document = Document()
        self._configure(document)
        self._set_running_furniture(document, account_name, period_label)
        self._add_cover(document, account_name, mode, period_label, len(articles), generated)
        current_year: str | None = None
        for article in articles:
            year = self._article_year(article)
            if year != current_year:
                current_year = year
                document.add_heading(f"{year} 年" if year.isdigit() else year, level=1)
            self._add_article(document, article)
        document.core_properties.title = self._document_title(account_name, mode)
        document.core_properties.subject = "微信公众号公开文章本地归档"
        document.core_properties.author = "微信公众号本地归档系统"
        document.save(temporary)
        self.verify(temporary, expected_articles=len(articles))
        temporary.replace(target)
        return target

    def export_summary(
        self,
        results: Sequence[Mapping[str, Any]],
        output_path: str | Path,
        *,
        cycle_key: str,
        generated_at: datetime | None = None,
    ) -> Path:
        target = Path(output_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_suffix(target.suffix + ".tmp")
        generated = generated_at or datetime.now().astimezone()
        document = Document()
        self._configure(document)
        self._set_running_furniture(document, "每周归档汇总", cycle_key)
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("微信公众号每周归档汇总")
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(31, 77, 120)
        self._set_east_asia(run, CJK_FONT)
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(f"周期：{cycle_key}  ｜  生成时间：{generated:%Y-%m-%d %H:%M}")
        document.add_heading("运行结果", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        self._set_table_geometry(table, (2700, 1500, 900, 4260))
        for cell, value in zip(table.rows[0].cells, ("公众号", "状态", "新增", "Word")):
            cell.text = value
            self._set_cell_fill(cell, "E8EEF5")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for result in results:
            cells = table.add_row().cells
            cells[0].text = str(result.get("account_name") or "")
            cells[1].text = str(result.get("status") or "")
            cells[2].text = str(result.get("exported_count") or 0)
            cells[3].text = str(result.get("word_path") or "无新增，不生成")
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        self._set_table_geometry(table, (2700, 1500, 900, 4260))
        document.save(temporary)
        self.verify(temporary)
        temporary.replace(target)
        return target

    def verify(self, path: str | Path, expected_articles: int | None = None) -> None:
        document = Document(path)
        if not document.paragraphs and not document.tables:
            raise ValueError(f"Word document is empty: {path}")
        if expected_articles is not None:
            headings = sum(
                1 for paragraph in document.paragraphs if paragraph.style.name == "Heading 2"
            )
            if headings != expected_articles:
                raise ValueError(
                    f"Word verification failed: expected {expected_articles} articles, found {headings}"
                )

    def _configure(self, document: Document) -> None:
        section = document.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        normal = document.styles["Normal"]
        normal.font.name = CJK_FONT
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(30, 30, 30)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.25
        for name, size, color, before, after in (
            ("Heading 1", 16, RGBColor(46, 116, 181), 18, 10),
            ("Heading 2", 13, RGBColor(46, 116, 181), 14, 7),
            ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
        ):
            style = document.styles[name]
            style.font.name = CJK_FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True

    def _add_cover(
        self,
        document: Document,
        account_name: str,
        mode: str,
        period_label: str,
        count: int,
        generated: datetime,
    ) -> None:
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_before = Pt(90)
        title.paragraph_format.space_after = Pt(8)
        run = title.add_run(self._document_title(account_name, mode))
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(31, 77, 120)
        self._set_east_asia(run, CJK_FONT)
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mode_label = "首次全量" if mode == "initial_full" else "本周增量"
        meta.add_run(
            f"{mode_label}  ｜  {count} 篇  ｜  {period_label}  ｜  {generated:%Y-%m-%d %H:%M}"
        )
        document.add_page_break()

    def _add_article(self, document: Document, article: Mapping[str, Any]) -> None:
        document.add_heading(str(self._value(article, "title") or "未命名文章"), level=2)
        meta = document.add_paragraph()
        meta.paragraph_format.space_after = Pt(5)
        label = self._publish_label(self._value(article, "publish_time"))
        author = str(self._value(article, "author") or "未注明")
        source_url = str(self._value(article, "source_url") or "")
        run = meta.add_run(f"{label}  ｜  作者：{author}  ｜  原文：")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(96, 96, 96)
        if source_url:
            self._add_hyperlink(meta, source_url, source_url)
        summary = self._clean_text(str(self._value(article, "summary") or ""))
        if summary:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.12)
            summary_run = paragraph.add_run(f"摘要：{summary}")
            summary_run.italic = True
            summary_run.font.size = Pt(9.5)
            summary_run.font.color.rgb = RGBColor(96, 96, 96)
        blocks = list(self._body_blocks(str(self._value(article, "content_text") or "")))
        if not blocks:
            blocks = ["[正文为空，原始链接已保留]" ]
        for block in blocks:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.add_run(block)

    def _set_running_furniture(self, document: Document, label: str, period_label: str) -> None:
        section = document.sections[0]
        header = section.header.paragraphs[0]
        header.text = f"{label}  |  {period_label}"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(112, 112, 112)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("第 ")
        run.font.size = Pt(8.5)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        footer.add_run(" 页").font.size = Pt(8.5)

    @staticmethod
    def _set_table_geometry(table: Any, widths: Sequence[int]) -> None:
        table.autofit = False
        properties = table._tbl.tblPr
        width = properties.first_child_found_in("w:tblW")
        if width is None:
            width = OxmlElement("w:tblW")
            properties.append(width)
        width.set(qn("w:w"), str(sum(widths)))
        width.set(qn("w:type"), "dxa")
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "120")
        indent.set(qn("w:type"), "dxa")
        properties.append(indent)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        properties.append(layout)
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for value in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(value))
            grid.append(column)
        for row in table.rows:
            for cell, value in zip(row.cells, widths):
                cell.width = value
                cell_properties = cell._tc.get_or_add_tcPr()
                cell_width = cell_properties.first_child_found_in("w:tcW")
                cell_width.set(qn("w:w"), str(value))
                cell_width.set(qn("w:type"), "dxa")
                margins = cell_properties.first_child_found_in("w:tcMar")
                if margins is None:
                    margins = OxmlElement("w:tcMar")
                    cell_properties.append(margins)
                for name, amount in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
                    item = OxmlElement(f"w:{name}")
                    item.set(qn("w:w"), str(amount))
                    item.set(qn("w:type"), "dxa")
                    margins.append(item)

    @staticmethod
    def _set_cell_fill(cell: Any, color: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        properties.append(shading)

    def _body_blocks(self, text: str) -> Iterable[str]:
        cleaned = text.replace("\u00a0", " ").replace("\r\n", "\n")
        pieces = re.split(r"\n\s*\n|(?<=。)\n", cleaned)
        for piece in pieces:
            lines = []
            for line in piece.splitlines():
                line = self._clean_text(line)
                if not line or line in NOISE_LINES or "轻点两下取消" in line:
                    continue
                if line.startswith("微信扫一扫可打开此内容"):
                    continue
                lines.append(line)
            block = "\n".join(lines).strip()
            if block:
                yield block

    @staticmethod
    def _clean_text(value: str) -> str:
        return re.sub(r"[ \t]+", " ", value.replace("\u00a0", " ")).strip()

    @staticmethod
    def _article_year(article: Mapping[str, Any]) -> str:
        value = str(IncrementalWordExporter._value(article, "publish_time") or "")
        match = re.match(r"(\d{4})-", value)
        return match.group(1) if match else "未注明日期"

    @staticmethod
    def _value(article: Mapping[str, Any], key: str) -> Any:
        try:
            return article[key]
        except (KeyError, IndexError, TypeError):
            getter = getattr(article, "get", None)
            return getter(key) if getter else None

    @staticmethod
    def _publish_label(value: Any) -> str:
        raw = str(value or "")
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError:
            return raw or "未注明日期"
        return f"{parsed.year}年{parsed.month}月{parsed.day}日 {parsed:%H:%M}"

    @staticmethod
    def _article_period_label(
        articles: Sequence[Mapping[str, Any]],
        fallback: str,
    ) -> str:
        dates: list[str] = []
        for article in articles:
            value = str(IncrementalWordExporter._value(article, "publish_time") or "")
            match = re.match(r"(\d{4}-\d{2}-\d{2})", value)
            if match:
                dates.append(match.group(1))
        if not dates:
            return fallback
        start = min(dates)
        end = max(dates)
        return start if start == end else f"{start} 至 {end}"

    @staticmethod
    def _document_title(account_name: str, mode: str) -> str:
        suffix = "公众号首次全量归档" if mode == "initial_full" else "公众号每周增量归档"
        return f"{account_name}\n{suffix}"

    @staticmethod
    def _set_east_asia(run: Any, font_name: str) -> None:
        run.font.name = font_name
        properties = run._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.append(fonts)
        fonts.set(qn("w:eastAsia"), font_name)
        fonts.set(qn("w:ascii"), font_name)
        fonts.set(qn("w:hAnsi"), font_name)
        fonts.set(qn("w:cs"), font_name)

    @staticmethod
    def _add_hyperlink(paragraph: Any, text: str, url: str) -> None:
        relation_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relation_id)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "2E74B5")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.append(color)
        properties.append(underline)
        run.append(properties)
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
