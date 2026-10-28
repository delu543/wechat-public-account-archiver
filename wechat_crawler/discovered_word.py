"""Faithful text-only Word for explicitly accepted, incomplete discoveries.

Reuse the existing Word layout and link primitives; never filter short bodies or
replace them with previews. Structural checks are not rendered-page acceptance.
"""
from __future__ import annotations

import html
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

from .biz_push_reader import permanent_url, public_identity
from .cache_delivery_receipts import file_sha256
from .crawler.dedup import normalized_text_hash
from .incremental_word import IncrementalWordExporter
from .passive_weekly_plan import publication_day

ANCHOR = re.compile(r'<a\s+class="normal_text_link mp_article_text_link"[^>]*href="([^"]+)"[^>]*>([^<]*)</a>')
BODY_STYLE = "Archive Body"


def visible_text(text):
    return ANCHOR.sub(lambda m: html.unescape(m.group(2)), text)


def compact(text):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))


class DiscoveredWordExporter(IncrementalWordExporter):
    def author(self, account, path):
        if not account["articles"]:
            raise ValueError("empty_account_word_forbidden")
        doc = Document()
        self._configure(doc)
        for style in doc.styles:
            for border in style._element.xpath('.//w:pBdr'):
                border.getparent().remove(border)
            for fonts in style._element.xpath('.//w:rFonts'):
                for attribute in list(fonts.attrib):
                    if "theme" in attribute.lower():
                        del fonts.attrib[attribute]
        for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
            doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
            doc.styles[name].font.name = "PingFang SC"
            doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "PingFang SC")
        body = doc.styles.add_style(BODY_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        body.base_style = doc.styles["Normal"]
        body.paragraph_format.widow_control = True
        self._set_running_furniture(doc, account["account_name"], "已发现增量")
        for section in doc.sections:
            for paragraph in section.header.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
        doc.add_paragraph(account["account_name"], style="Title")
        doc.add_paragraph("已发现增量", style="Subtitle")
        doc.add_paragraph(f"检索区间：{account['since']} 至 {account['until']}（北京时间）")
        days = [publication_day(a["publish_time"]) for a in account["articles"]]
        doc.add_paragraph(f"本册 {len(days)} 篇，含短内容；实际发表日期：{min(days)} 至 {max(days)}。")
        doc.add_paragraph("仅收录已取得且通过本地核查的正文文字及链接，不代表区间内全部发表内容已齐全。"
                          "图片、音视频请打开原文查看。完整归档节点未推进，未取得的内容仍保留为缺口。")
        for gap in account["gaps"]:
            interval = gap.get("range", {})
            detail = (f"；{interval.get('since') or interval.get('since_date')} 至 "
                      f"{interval.get('until') or interval.get('until_date')}" if interval else "")
            reason = {"known_preview_not_recovered": "仍为旧片段，未收录",
                      "historical_interval_unverified": "历史区间覆盖尚未核实",
                      "body_unverified": "正文尚未核实，未收录",
                      "conflicting_versions": "存在不同版本，未收录",
                      "previously_delivered_body_changed": "已交付原文发生变化，保留旧版待核对"}.get(
                          gap["reason"], "来源或正文尚待核实")
            doc.add_paragraph(f"已知缺口：{gap.get('title') or '区间覆盖'}（{reason}{detail}）")
        for index, article in enumerate(account["articles"]):
            if normalized_text_hash(article["content_text"]) != article["content_hash"]:
                raise ValueError("staged_body_changed")
            heading = doc.add_heading(article["title"], level=2)
            if index == 0:
                heading.paragraph_format.page_break_before = True
            meta = doc.add_paragraph(self._publish_label(article["publish_time"]) + "  ·  ")
            meta.paragraph_format.keep_with_next = True
            self._add_hyperlink(meta, "打开原文", article["url"])
            if article["title_kind"] != "original_title":
                doc.add_paragraph("标题取自正文首行，原内容没有独立标题。")
            for line in article["content_text"].splitlines():
                if not line.strip():
                    continue
                paragraph, previous = doc.add_paragraph(style=BODY_STYLE), 0
                for match in ANCHOR.finditer(line):
                    paragraph.add_run(line[previous:match.start()])
                    label = html.unescape(match.group(2))
                    key = public_identity(html.unescape(match.group(1)))
                    if key:
                        self._add_hyperlink(paragraph, label, permanent_url(key))
                    else:
                        paragraph.add_run(label)  # No transient/private URL is emitted.
                    previous = match.end()
                paragraph.add_run(line[previous:])
        doc.core_properties.title = account["account_name"] + "｜已发现增量"
        doc.core_properties.author = "微信公众号本地归档"
        doc.core_properties.language = "zh-CN"
        with Path(path).open("xb") as handle:
            Path(path).chmod(0o600)
            doc.save(handle)
        return verify_discovered_word(path, account)


def verify_discovered_word(path, account):
    doc = Document(path)
    headings, bodies, current = [], [], None
    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Heading 2":
            headings.append(paragraph.text)
            current = []
            bodies.append(current)
        elif paragraph.style.name == BODY_STYLE:
            if current is None:
                raise ValueError("word_body_without_article")
            current.append("".join(paragraph._p.xpath('.//w:t/text()')))
    if headings != [a["title"] for a in account["articles"]] or len(bodies) != len(account["articles"]):
        raise ValueError("word_membership_mismatch")
    for article, parts in zip(account["articles"], bodies):
        if compact("\n".join(parts)) != compact(visible_text(article["content_text"])):
            raise ValueError("word_body_text_mismatch")
    links = {r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK and r.is_external}
    if not {a["url"] for a in account["articles"]} <= links:
        raise ValueError("word_source_links_missing")
    for url in links:
        key = public_identity(url)
        if not key or permanent_url(key) != url:
            raise ValueError("word_nonsanitized_link")
    if "已发现增量" not in "\n".join(p.text for p in doc.paragraphs):
        raise ValueError("word_coverage_label_missing")
    return {"heading_count": len(headings), "body_text_conserved": True, "sha256": file_sha256(path),
            "render_reviewed": False}
