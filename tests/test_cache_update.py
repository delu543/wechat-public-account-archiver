"""Synthetic-only repeatable delivery; never reads the live client or archive."""
import copy
import json
from pathlib import Path

from docx import Document
import pytest

from test_passive_weekly import prepare, url
from wechat_crawler import cache_update, main as cli
from wechat_crawler.cache_capture_all import prepare_scopes
from wechat_crawler.cache_capture_store import CaptureStore
from wechat_crawler.cache_capture_validation import record_value
from wechat_crawler.cache_delivery_receipts import delivered_members, file_sha256, read_delivery
from wechat_crawler.cache_update_export import confirm_delivery, output_lock, write_once
from wechat_crawler.crawler.storage import Storage
from wechat_crawler.discovered_export_plan import select_discovered
from wechat_crawler.discovered_word import verify_discovered_word


def arguments(*values):
    return cli.build_parser().parse_args(["cache-update", "--until", "2026-10-28", *values])


def setup(tmp_path, *, existing=True, accounts=1):
    storage, automation, settings = prepare(tmp_path, paused=True, accounts=accounts)
    config = {"database": {"path": str(storage.db_path)}, "output": {"path": str(tmp_path / "output")},
              "passive_weekly": settings}
    root = tmp_path / "output" / "cache-updates" / "stages"
    scopes, paths = prepare_scopes(config, "2026-10-28", root)
    records = [dict(biz=scopes[0]["biz"], wechat_id=scopes[0]["wechat_id"], account_name=scopes[0]["account_name"],
                    url=url(mid=i), title=f"文章{i}", title_kind="original_title",
                    publish_time=f"2026-10-{i+10:02d}T12:00:00+08:00", content_text=f"正文{i}\n允许\n分析",
                    body_ready=True) for i in (2, 3)]
    if existing:
        append(paths[scopes[0]["biz"]], scopes[0], records)
    return config, scopes, paths, records, storage


def append(path, scope, records):
    store = CaptureStore(path, scope)
    try:
        store.ingest(dict(status="read", records=records, read_errors=0, unsupported_values=0,
                          metadata_conflicts=0, scan_within_bounds=True))
    finally:
        store.close()


def export(config, *options):
    return cache_update.run(arguments("--export", "--accept-discovered", *options), config)["export"]


def reviewed(config):
    event = export(config)
    # Test the operator attestation state transition, not real rendered QA.
    confirm_delivery(event["folder"], render_reviewed=True)
    return Path(event["folder"])


def test_default_only_reads_retained_evidence_and_reports_unknown_coverage(tmp_path):
    config, _, paths, _, storage = setup(tmp_path)
    before = {p: p.read_bytes() for p in [storage.db_path, *paths.values()]}
    event = cache_update.run(arguments(), config, source=lambda _: pytest.fail("no live read"))
    assert event["accounts"][0]["export_count"] == 2
    assert not event["accounts"][0]["full_interval_verified"]
    assert event["pending_batches"] == []
    assert all(p.read_bytes() == value for p, value in before.items())


def test_cli_dispatches_before_product_init(tmp_path, monkeypatch, capsys):
    config, *_ = setup(tmp_path)
    monkeypatch.setattr(cli, "load_config", lambda _: config)
    monkeypatch.setattr(Storage, "init_db", lambda *_: pytest.fail("product init forbidden"))
    assert cli.main(["cache-update", "--check", "--until", "2026-10-28"]) == 0
    assert json.loads(capsys.readouterr().out)["product_writes"] == 0


def test_missing_account_is_unknown_and_does_not_block_available_account(tmp_path):
    config, *_ = setup(tmp_path, accounts=2)
    event = cache_update.run(arguments(), config)
    assert event["accounts"][1]["retained_count"] is None
    assert not event["accounts"][1]["zero_increment_claimed"]
    result = export(config)
    receipt = json.loads((Path(result["folder"]) / "delivery-manifest.json").read_text())
    assert result["documents"] == 1 and len(receipt["accounts_without_document"]) == 1


def test_no_empty_word_or_delivery_batch(tmp_path):
    config, *_ = setup(tmp_path, existing=False)
    result = export(config)
    assert result["status"] == "no_eligible_retained_articles"
    assert not list(Path(config["output"]["path"]).rglob("*.docx"))


@pytest.mark.parametrize("flags", [("--export",), ("--check", "--collect"), ("--check", "--export"),
                                  ("--collect", "--seconds", "1801")])
def test_invalid_action_stops_without_live_io(tmp_path, flags):
    config, *_ = setup(tmp_path, existing=False)
    with pytest.raises(ValueError):
        cache_update.run(arguments(*flags), config, source=lambda _: pytest.fail("no read"))
    assert not Path(config["output"]["path"]).exists()


def test_paused_collection_stops_before_new_folders_or_live_read(tmp_path):
    config, *_ = setup(tmp_path, existing=False)
    event = cache_update.run(arguments("--collect"), config, source=lambda _: pytest.fail("no read"))
    assert event["status"] == "paused" and not Path(config["output"]["path"]).exists()


def test_one_shared_collection_keeps_pause_and_owned_product_unchanged(tmp_path):
    config, _, _, records, storage = setup(tmp_path, existing=False)
    before = storage.db_path.read_bytes()
    calls = []
    def source(scopes):
        calls.append(scopes)
        return dict(status="read", records=records, read_errors=0, unsupported_values=0,
                    metadata_conflicts=0, scan_within_bounds=True)
    event = cache_update.run(arguments("--collect", "--seconds", "1", "--allow-local-stage-while-paused"), config, source=source)
    assert len(calls) == 1 and event["capture"]["shared_source_reads"] == 1
    assert event["accounts"][0]["export_count"] == 2
    assert storage.db_path.read_bytes() == before and event["automation_paused"]


def test_collection_observes_pause_change_before_write(tmp_path, monkeypatch):
    config, _, _, records, _ = setup(tmp_path, existing=False)
    state = dict(paused=True, updated_at="old")
    monkeypatch.setattr(cache_update, "read_control", lambda _: state)
    def source(_):
        state["updated_at"] = "new"
        return dict(status="read", records=records, read_errors=0)
    event = cache_update.run(arguments("--collect", "--seconds", "1", "--allow-local-stage-while-paused"), config, source=source)
    assert event["status"] == "control_changed"
    assert not list(Path(config["output"]["path"]).rglob("*.sqlite"))


def test_next_day_uses_retained_evidence_without_claiming_refresh(tmp_path):
    config, _, paths, _, _ = setup(tmp_path)
    before = {p:p.read_bytes() for p in paths.values()}
    args = arguments()
    args.until = "2026-10-29"
    event = cache_update.run(args, config)
    assert event["accounts"][0]["retained_until"] == "2026-10-28"
    args.collect, args.allow_local_stage_while_paused, args.seconds = True, True, 1
    event = cache_update.run(args, config, source=lambda _: dict(status="read", records=[], read_errors=0))
    assert event["accounts"][0]["export_count"] == 2
    assert all(p.read_bytes() == value for p,value in before.items())


def test_exclusions_persist_and_explicit_conflicting_selection_is_not_silent(tmp_path):
    config, *_ = setup(tmp_path, accounts=2)
    config["cache_update"] = {"excluded_biz": ["other-account"]}
    event = cache_update.run(arguments(), config)
    assert len(event["accounts"]) == 1 and event["excluded_biz"] == ["other-account"]
    with pytest.raises(ValueError, match="selected_account_is_excluded"):
        cache_update.run(arguments("--account-biz", "other-account"), config)
    with pytest.raises(ValueError):
        cache_update.run(arguments("--account-biz", "not-registered"), config)


def test_export_review_then_dedup_without_advancing_product(tmp_path):
    config, _, _, _, storage = setup(tmp_path)
    before = storage.db_path.read_bytes()
    result = export(config)
    folder = Path(result["folder"])
    assert result["articles"] == 2 and result["status"] == "awaiting_render_review"
    assert cache_update.run(arguments(), config)["accounts"][0]["export_count"] == 2
    with pytest.raises(ValueError, match="render_review_required"):
        confirm_delivery(folder)
    confirm_delivery(folder, render_reviewed=True)
    event = cache_update.run(arguments(), config)
    assert event["accounts"][0]["export_count"] == 0
    assert event["accounts"][0]["delivered_overlap_count"] == 2
    assert event["accounts"][0]["suggested_refresh_through"] == "2026-10-13"
    assert export(config)["documents"] == 0
    assert storage.db_path.read_bytes() == before


def test_pending_batch_not_lost_when_new_articles_arrive(tmp_path):
    config, scopes, paths, records, _ = setup(tmp_path)
    first = export(config)
    new = dict(records[0], url=url(mid=4), title="稍后到达", content_text="新来正文", publish_time="2026-10-24T12:00:00+08:00")
    append(paths[scopes[0]["biz"]], scopes[0], [new])
    second = export(config)
    assert second["folder"] == first["folder"] and second["resumed"]
    confirm_delivery(first["folder"], render_reviewed=True)
    third = export(config)
    assert third["articles"] == 1 and third["folder"] != first["folder"]


def test_interrupted_after_plan_resumes_without_replacing_old_word(tmp_path, monkeypatch):
    config, *_ = setup(tmp_path)
    from wechat_crawler import cache_update_export as module
    original = module.DiscoveredWordExporter.author
    def interrupted(self, account, path):
        original(self, account, path)
        raise KeyboardInterrupt()
    monkeypatch.setattr(module.DiscoveredWordExporter, "author", interrupted)
    with pytest.raises(KeyboardInterrupt):
        export(config)
    word = next(Path(config["output"]["path"]).rglob("*.docx"))
    content = word.read_bytes()
    monkeypatch.setattr(module.DiscoveredWordExporter, "author", lambda *_: pytest.fail("do not reauthor"))
    assert export(config)["resumed"]
    assert word.read_bytes() == content


@pytest.mark.parametrize("change", ["word", "manifest", "review", "cross_account", "temporary_url", "duplicate_member"])
def test_changed_or_invalid_receipt_is_never_used_to_skip_articles(tmp_path, change):
    config, *_ = setup(tmp_path)
    folder = reviewed(config)
    manifest_path = folder / "delivery-manifest.json"
    manifest = json.loads(manifest_path.read_text())
    if change == "word":
        (folder / manifest["artifacts"][0]["file"]).write_bytes(b"corrupt")
    elif change == "review":
        (folder / "verification.json").write_text("{}")
    elif change == "manifest":
        manifest_path.write_text("{}")
    else:
        member = manifest["artifacts"][0]["members"][0]
        if change == "cross_account": member["url"] = url("foreign-account", 2)
        elif change == "temporary_url": member["url"] += "&key=DO_NOT_EXPORT"
        else: manifest["artifacts"][0]["members"][1] = member
        manifest_path.write_text(json.dumps(manifest))
        review = json.loads((folder / "verification.json").read_text())
        review["output_manifest_sha256"] = file_sha256(manifest_path)
        (folder / "verification.json").write_text(json.dumps(review))
    with pytest.raises(ValueError):
        cache_update.run(arguments(), config)


def test_short_body_reserved_words_and_embedded_link_conserved(tmp_path):
    config, scopes, paths, records, _ = setup(tmp_path, existing=False)
    body = f'允许\n分析\n2 < 3\n<a class="normal_text_link mp_article_text_link" href="{url(mid=9)}&key=secret">原文片段</a>\n🙂'
    row = dict(records[0], content_text=body, title_kind="first_line")
    append(paths[scopes[0]["biz"]], scopes[0], [row])
    folder = Path(export(config)["folder"])
    document = Document(next(folder.glob("*.docx")))
    texts = "\n".join("".join(p._p.xpath('.//w:t/text()')) for p in document.paragraphs)
    assert all(t in texts for t in ("允许", "分析", "2 < 3", "原文片段", "🙂", "没有独立标题"))
    assert "key=" not in str([r.target_ref for r in document.part.rels.values()])


def test_changed_delivered_body_is_gap_not_replacement(tmp_path):
    config, scopes, _, records, _ = setup(tmp_path)
    folder = reviewed(config)
    receipts = delivered_members([folder / "delivery-manifest.json"])
    account = dict(_articles=[], incomplete_article_ids=[])
    changed = record_value(dict(records[0], content_text="正文已改版"), scopes[0])
    result = select_discovered(scopes[0], account, [changed], config["passive_weekly"], delivered=receipts[scopes[0]["biz"]])
    assert result["export_count"] == 0 and result["gaps"][0]["reason"] == "previously_delivered_body_changed"


def test_existing_word_edit_stops_pending_resume_without_overwrite(tmp_path):
    config, *_ = setup(tmp_path)
    folder = Path(export(config)["folder"])
    word = next(folder.glob("*.docx"))
    word.write_bytes(b"user-edited-evidence")
    with pytest.raises(ValueError):
        export(config)
    assert word.read_bytes() == b"user-edited-evidence"


def test_same_body_new_url_suppressed_by_verified_receipt(tmp_path):
    config, scopes, paths, records, _ = setup(tmp_path)
    reviewed(config)
    append(paths[scopes[0]["biz"]], scopes[0], [dict(records[0], url=url(mid=99))])
    result = cache_update.run(arguments(), config)
    assert result["accounts"][0]["export_count"] == 0
    assert result["accounts"][0]["delivered_overlap_count"] == 3


def test_receipt_does_not_clear_historical_coverage_gap(tmp_path):
    config, scopes, _, records, _ = setup(tmp_path)
    folder = reviewed(config)
    account = dict(_articles=[], incomplete_article_ids=[], recovery_gap={"since":"2026-09-23","until":"2026-09-28"})
    result = select_discovered(scopes[0], account, [record_value(r, scopes[0]) for r in records],
        config["passive_weekly"], delivered=delivered_members([folder / "delivery-manifest.json"])[scopes[0]["biz"]])
    assert result["export_count"] == 0
    assert result["gaps"][0]["range"] == account["recovery_gap"]
    assert not result["product_checkpoint_advanced"]


def test_lock_prevents_concurrent_delivery(tmp_path):
    with output_lock(tmp_path):
        with pytest.raises(BlockingIOError):
            with output_lock(tmp_path):
                pytest.fail("lock did not serialize")


def test_cover_explains_historical_gap_in_plain_language(tmp_path):
    config, *_ = setup(tmp_path)
    config["passive_weekly"]["recovery_gaps"] = {"biz-account": {"since":"2026-09-23","until":"2026-09-28"}}
    folder = Path(export(config)["folder"])
    doc = Document(next(folder.glob("*.docx")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "历史区间覆盖尚未核实；2026-09-23 至 2026-09-28" in text
    assert "historical_interval_unverified" not in text and "'since'" not in text
    assert all(str(r.font.color.rgb) == "000000" for p in doc.sections[0].header.paragraphs for r in p.runs)
