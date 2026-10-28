"""Cross-run acceptance, including new source rows during an interruption."""
from pathlib import Path

from docx import Document

from test_passive_weekly import (Source, _distinct_fetcher, _three_records_source,
                                prepare, runner)
from wechat_crawler.passive_weekly_store import PassiveWeeklyStore


def test_interrupted_word_resumes_with_newly_arrived_articles(tmp_path, monkeypatch):
    storage, _, settings = prepare(tmp_path)
    original = PassiveWeeklyStore.finish
    monkeypatch.setattr(PassiveWeeklyStore, "finish", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("interrupted")))
    first = runner(tmp_path, storage, settings, sources=Source(), fetcher=_distinct_fetcher()).run("2026-10-26")
    old_path = Path(first["accounts"][0]["word_path"])
    assert old_path.exists()
    monkeypatch.setattr(PassiveWeeklyStore, "finish", original)
    second = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher()).run("2026-10-26")
    assert second["status"] == "completed"
    with storage.connect() as conn:
        pending = conn.execute("SELECT count(*) FROM export_batches WHERE status='preparing'").fetchone()[0]
        assert pending == 0
    document = Document(second["accounts"][0]["word_path"])
    assert {p.text for p in document.paragraphs if p.style.name == "Heading 2"} == {"新文章2", "新文章3"}
    assert old_path.exists()  # never overwrite/delete the interrupted artifact


def test_summary_is_refreshed_after_recovery_without_overwriting_old_summary(tmp_path):
    storage, _, settings = prepare(tmp_path)
    first = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher(fail_second=True)).run("2026-10-26", allow_partial=True)
    old_summary = Path(first["summary_path"])
    before = old_summary.read_bytes()
    second = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher()).run("2026-10-26")
    assert second["status"] == "completed"
    assert second["summary_path"] != first["summary_path"]
    assert old_summary.read_bytes() == before
    row = Document(second["summary_path"]).tables[0].rows[1]
    assert row.cells[1].text == "已完成"
    assert Path(second["accounts"][0]["word_path"]).name in row.cells[3].text


def test_repeated_completed_cycle_reuses_summary_and_does_not_export_again(tmp_path):
    storage, _, settings = prepare(tmp_path)
    task = runner(tmp_path, storage, settings, sources=Source(), fetcher=_distinct_fetcher())
    first = task.run("2026-10-26")
    second = task.run("2026-10-26")
    assert second["accounts"][0]["exported_count"] == 0
    assert second["summary_path"] == first["summary_path"]
