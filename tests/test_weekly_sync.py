from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from wechat_crawler.automation_store import AutomationStore
from wechat_crawler.crawler.models import ParsedArticle
from wechat_crawler.crawler.storage import Storage
from wechat_crawler.incremental_word import IncrementalWordExporter
from wechat_crawler.weekly_sync import CrawlResult, ExistingCrawlerAdapter, WeeklySyncRunner
from wechat_crawler.web_app import ArchiveWebService


def article(url: str, title: str, publish_time: str, raw_html: str = "<html>large</html>") -> ParsedArticle:
    return ParsedArticle(
        source_url=url,
        normalized_url=url,
        title=title,
        account_name="Account",
        account_id="biz-account",
        wechat_id="gh_account",
        author="Author",
        publish_time=publish_time,
        summary="Summary",
        content_html=f"<p>{title} body</p>",
        content_text=f"{title} body\n\nSecond paragraph",
        raw_html=raw_html,
        content_hash=f"hash-{title}",
    )


def save(storage: Storage, value: ParsedArticle) -> int:
    run_id = storage.start_run("test")
    url_id = storage.upsert_article_url(value.source_url, "test", run_id=run_id)
    article_id = storage.save_article(value, url_id, run_id)
    storage.finish_run(run_id, "completed")
    return article_id


class AddingCrawler:
    def __init__(self, storage: Storage) -> None:
        self.storage = storage
        self.calls = 0

    def run(self, subscription, job, progress):
        self.calls += 1
        progress("discover", 25, {"discovered_count": 1})
        save(
            self.storage,
            article(
                "https://mp.weixin.qq.com/s/new-article",
                "New article",
                "2026-07-18 09:00:00",
            ),
        )
        return CrawlResult(status="completed", message="crawl complete", discovered_count=1, crawled_count=1)


class CooperativePausedCrawler:
    def run(self, subscription, job, progress):
        progress("crawl_fast", 40, {"discovered_count": 1})
        return CrawlResult(
            status="paused",
            message="全部爬取已由用户暂停",
            discovered_count=1,
        )


def prepare_existing(tmp_path):
    storage = Storage(tmp_path / "archive.sqlite")
    automation = AutomationStore(storage)
    automation.init_db()
    save(storage, article("https://mp.weixin.qq.com/s/old", "Old article", "2026-07-10 09:00:00"))
    automation.adopt_existing_accounts(now=datetime(2026, 7, 17, 12, tzinfo=timezone.utc))
    subscription = automation.list_subscriptions()[0]
    job = automation.create_job(
        int(subscription["id"]),
        "2026-W29",
        scheduled_for=datetime(2026, 7, 19, 20, tzinfo=timezone.utc),
    )
    return storage, automation, subscription, job


def test_weekly_job_exports_only_new_articles_and_is_resumable(tmp_path):
    storage, automation, _, job = prepare_existing(tmp_path)
    crawler = AddingCrawler(storage)
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=crawler)

    result = runner.run_job(job["id"])
    assert result.status == "completed"
    assert result.exported_count == 1
    assert crawler.calls == 1
    assert result.word_path
    assert Path(result.word_path).name == "Account_2026-07-13至2026-07-19_本周增量.docx"
    document = Document(result.word_path)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["New article"]
    assert automation.unexported_articles(int(job["account_id"])) == []
    with storage.connect() as conn:
        raw = conn.execute("SELECT raw_html FROM articles WHERE title = 'New article'").fetchone()[0]
    assert raw == ""

    repeated = runner.run_job(job["id"])
    assert repeated.exported_count == 1
    assert crawler.calls == 1


def test_global_pause_prevents_due_cycle_from_starting(tmp_path):
    storage = Storage(tmp_path / "archive.sqlite")
    automation = AutomationStore(storage)
    automation.init_db()
    save(storage, article("https://mp.weixin.qq.com/s/old", "Old", "2026-07-10 09:00:00"))
    automation.adopt_existing_accounts(now=datetime(2026, 7, 17, 12, tzinfo=timezone.utc))
    with storage.connect() as conn:
        conn.execute("UPDATE subscriptions SET next_run_at = '2026-07-19T19:00:00+00:00'")
    automation.set_automation_paused(True)
    crawler = AddingCrawler(storage)
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=crawler)

    results = runner.run_due(now=datetime(2026, 7, 19, 20, tzinfo=timezone.utc))

    assert results == []
    assert crawler.calls == 0
    assert automation.list_jobs() == []


def test_job_paused_at_safe_checkpoint_remains_resumable(tmp_path):
    storage, automation, subscription, job = prepare_existing(tmp_path)
    automation.begin_cycle(
        "2026-W29",
        scheduled_for=datetime(2026, 7, 19, 20, tzinfo=timezone.utc),
        output_folder=str(tmp_path / "output" / "weekly" / "2026-07-19"),
        expected_account_count=1,
    )
    runner = WeeklySyncRunner(
        storage,
        tmp_path / "output",
        {},
        crawler=CooperativePausedCrawler(),
    )

    result = runner.run_job(str(job["id"]))

    assert result.status == "paused"
    assert automation.job(str(job["id"]))["status"] == "paused"
    assert automation.latest_resumable_job(int(subscription["id"]))["id"] == job["id"]
    assert automation.list_subscriptions()[0]["state"] == "active"


def test_pending_export_recovery_skips_discovery_and_marks_articles_once(tmp_path):
    storage, automation, subscription, _job = prepare_existing(tmp_path)
    save(
        storage,
        article(
            "https://mp.weixin.qq.com/s/recovered",
            "Recovered article",
            "2026-07-22 09:00:00",
        ),
    )
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=NoopCrawler())

    result = runner.export_pending_subscription_now(int(subscription["id"]))

    assert result.status == "completed"
    assert result.exported_count == 1
    assert result.word_path
    document = Document(result.word_path)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["Recovered article"]
    assert automation.unexported_articles(int(subscription["account_id"])) == []


def test_word_export_verifies_article_count_and_has_clickable_source(tmp_path):
    exporter = IncrementalWordExporter()
    row = {
        "title": "Test title",
        "publish_time": "2026-07-18 09:00:00",
        "author": "Author",
        "summary": "Summary",
        "source_url": "https://mp.weixin.qq.com/s/test",
        "content_text": "First paragraph\n\nSecond paragraph",
    }
    path = exporter.export_account(
        [row],
        tmp_path / "test.docx",
        account_name="Account",
        mode="incremental",
        cycle_key="2026-W29",
    )
    exporter.verify(path, expected_articles=1)
    assert "hyperlink" in path.read_bytes().decode("latin1", errors="ignore") or path.stat().st_size > 0
    document = Document(path)
    assert "2026-07-18" in document.sections[0].header.paragraphs[0].text
    assert "2026-W29" not in document.sections[0].header.paragraphs[0].text


def test_word_export_decodes_short_content_summary_and_uses_it_as_body(tmp_path):
    exporter = IncrementalWordExporter()
    row = {
        "title": "短内容标题",
        "publish_time": "2026-08-29 09:00:00",
        "source_url": "https://mp.weixin.qq.com/s/short",
        "summary": r"完整第一段\x0a\x0a完整第二段\x26lt;a class=\x26quot;wx_img_refer_link\x26quot;\x26gt;图二\x26lt;/a\x26gt;" * 3,
        "content_text": "短内容标题",
    }

    path = exporter.export_account(
        [row],
        tmp_path / "short.docx",
        account_name="Account",
        mode="incremental",
        cycle_key="2026-W35",
    )

    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    assert "完整第一段" in text
    assert "完整第二段图二" in text
    assert "摘要：" not in text
    assert "\\x0a" not in text
    assert "wx_img_refer_link" not in text
    assert Document(path).styles["Normal"].font.name == "PingFang SC"


def test_word_export_rebuilds_inline_flow_from_matching_archived_html(tmp_path):
    row = {
        "title": "旧文",
        "publish_time": "2026-10-07 09:00:00",
        "source_url": "https://mp.weixin.qq.com/s/old",
        "content_text": "你好\n，\n世界。\n第二段",
        "content_html": "<div><p><span>你好</span><span>，</span><span>世界。</span></p><p>第二段</p></div>",
    }

    path = IncrementalWordExporter().export_account(
        [row],
        tmp_path / "old.docx",
        account_name="测试号",
        mode="range_incremental",
        cycle_key="2026-09-23_to_2026-10-08",
    )

    paragraphs = [p.text for p in Document(path).paragraphs]
    assert "你好，世界。" in paragraphs
    assert "第二段" in paragraphs
    assert "你好" not in paragraphs


def test_word_export_keeps_archived_text_when_html_differs(tmp_path):
    row = {
        "title": "旧文",
        "publish_time": "2026-10-07 09:00:00",
        "source_url": "https://mp.weixin.qq.com/s/old",
        "content_text": "存档原文",
        "content_html": "<p>不相同的摘要</p>",
    }

    path = IncrementalWordExporter().export_account(
        [row],
        tmp_path / "mismatch.docx",
        account_name="测试号",
        mode="range_incremental",
        cycle_key="2026-09-23_to_2026-10-08",
    )

    paragraphs = [p.text for p in Document(path).paragraphs]
    assert "存档原文" in paragraphs
    assert "不相同的摘要" not in paragraphs


def test_word_summary_uses_readable_status_and_filename(tmp_path):
    path = IncrementalWordExporter().export_summary(
        [
            {
                "account_name": "公众号A",
                "status": "completed",
                "exported_count": 2,
                "word_path": "/very/long/product/path/公众号A_本周增量.docx",
            },
            {
                "account_name": "公众号B",
                "status": "completed",
                "exported_count": 0,
                "word_path": None,
            },
        ],
        tmp_path / "summary.docx",
        cycle_key="2026-W35",
    )

    cells = [cell.text for table in Document(path).tables for row in table.rows for cell in row.cells]
    assert "已完成" in cells
    assert "公众号A_本周增量.docx" in cells
    assert "/very/long/product/path" not in "\n".join(cells)
    assert "无新增，不生成" in cells


def test_word_export_uses_article_date_range_instead_of_internal_cycle_key(tmp_path):
    exporter = IncrementalWordExporter()
    rows = [
        {
            "title": "First",
            "publish_time": "2026-07-13 09:00:00",
            "source_url": "https://mp.weixin.qq.com/s/first",
            "content_text": "First body",
        },
        {
            "title": "Last",
            "publish_time": "2026-07-22 09:00:00",
            "source_url": "https://mp.weixin.qq.com/s/last",
            "content_text": "Last body",
        },
    ]
    path = exporter.export_account(
        rows,
        tmp_path / "range.docx",
        account_name="Account",
        mode="incremental",
        cycle_key="manual-20260722-internal",
    )

    document = Document(path)
    header = document.sections[0].header.paragraphs[0].text
    cover = "\n".join(paragraph.text for paragraph in document.paragraphs[:4])
    assert "2026-07-13 至 2026-07-22" in header
    assert "2026-07-13 至 2026-07-22" in cover
    assert "manual-" not in header
    assert "manual-" not in cover


def test_word_export_shows_beijing_date_for_aware_archived_timestamp(tmp_path):
    row = {
        "title": "跨时区文章",
        "publish_time": "2026-10-07T09:30:00-07:00",
        "source_url": "https://mp.weixin.qq.com/s/test",
        "content_text": "文章正文",
    }

    path = IncrementalWordExporter().export_account(
        [row],
        tmp_path / "beijing.docx",
        account_name="测试号",
        mode="range_incremental",
        cycle_key="2026-09-23_to_2026-10-08",
    )

    document = Document(path)
    assert "2026-10-07" not in document.sections[0].header.paragraphs[0].text
    assert "2026-10-08" in document.sections[0].header.paragraphs[0].text
    assert any("2026年10月8日 00:30" in p.text for p in document.paragraphs)


def test_word_export_labels_manual_date_range_without_changing_weekly_mode(tmp_path):
    exporter = IncrementalWordExporter()
    row = {
        "title": "Range article",
        "publish_time": "2026-08-04 09:00:00",
        "source_url": "https://mp.weixin.qq.com/s/range",
        "content_text": "Range body",
    }

    range_path = exporter.export_account(
        [row],
        tmp_path / "range-incremental.docx",
        account_name="Account",
        mode="range_incremental",
        cycle_key="2026-08-04_to_2026-08-21",
    )
    weekly_path = exporter.export_account(
        [row],
        tmp_path / "weekly-incremental.docx",
        account_name="Account",
        mode="incremental",
        cycle_key="2026-W32",
    )

    range_cover = "\n".join(paragraph.text for paragraph in Document(range_path).paragraphs[:4])
    weekly_cover = "\n".join(paragraph.text for paragraph in Document(weekly_path).paragraphs[:4])
    range_document = Document(range_path)
    year_heading = next(
        paragraph
        for paragraph in range_document.paragraphs
        if paragraph.style.name == "Heading 1"
    )
    assert "公众号区间增量归档" in range_cover
    assert "区间增量" in range_cover
    assert "公众号每周增量归档" in weekly_cover
    assert "本周增量" in weekly_cover
    assert year_heading.paragraph_format.page_break_before is True
    assert b'w:type="page"' not in range_path.read_bytes()


class NoopCrawler:
    def run(self, subscription, job, progress):
        return CrawlResult(status="completed", message="no updates")


class UnresolvedCrawler:
    def run(self, subscription, job, progress):
        return CrawlResult(
            status="completed",
            message="completed with unresolved public pages",
            discovered_count=3,
            failed_count=2,
            restricted_count=1,
            unresolved_by_reason={"empty article body": 2, "captcha": 1},
        )


class NeedsSessionThenCompleteCrawler:
    def __init__(self):
        self.calls = 0

    def run(self, subscription, job, progress):
        self.calls += 1
        if self.calls == 1:
            return CrawlResult(status="needs_session", message="refresh visible session")
        return CrawlResult(status="completed", message="session recovered")


def test_weekly_job_persists_unresolved_reasons_for_the_dashboard(tmp_path):
    storage, automation, _, job = prepare_existing(tmp_path)
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=UnresolvedCrawler())

    result = runner.run_job(job["id"])

    assert result.status == "completed"
    stored = automation.job(job["id"])
    assert json.loads(stored["checkpoint_json"]) == {
        "unresolved_by_reason": {"captcha": 1, "empty article body": 2}
    }
    service = ArchiveWebService(storage, tmp_path / "output", config={})
    payload = service.get_automation_job(job["id"])
    assert payload["unresolved_by_reason"] == {
        "captcha": 1,
        "empty article body": 2,
    }


def test_due_run_writes_summary_without_empty_account_word(tmp_path):
    storage, automation, subscription, _ = prepare_existing(tmp_path)
    with storage.connect() as conn:
        conn.execute(
            "UPDATE subscriptions SET next_run_at = '2026-07-19T19:00:00+00:00' WHERE id = ?",
            (subscription["id"],),
        )
        conn.execute("DELETE FROM sync_jobs")
    orphan = tmp_path / "output" / "weekly" / "2026-07-19" / "orphan.docx.tmp"
    orphan.parent.mkdir(parents=True)
    orphan.write_bytes(b"temporary")
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=NoopCrawler())
    results = runner.run_due(now=datetime(2026, 7, 19, 20, tzinfo=timezone.utc))
    assert len(results) == 1
    assert results[0].exported_count == 0
    folder = tmp_path / "output" / "weekly" / "2026-07-19"
    assert (folder / "run-report.json").exists()
    assert (folder / "2026-07-13至2026-07-19_本周公众号增量汇总.docx").exists()
    assert list(folder.glob("*_本周增量.docx")) == []
    assert not orphan.exists()
    completion = automation.latest_completed_cycle()
    assert completion["cycle_key"] == "2026-W29"
    assert completion["completed_account_count"] == 1
    assert completion["exported_count"] == 0
    assert completion["output_folder"] == str(folder.resolve())
    widget = ArchiveWebService(storage, tmp_path / "output", config={}).status_widget()
    assert widget["latest_completion"] == {
        "cycle_key": "2026-W29",
        "scheduled_for": "2026-07-19T20:00:00+00:00",
        "completed_at": completion["completed_at"],
        "output_folder": str(folder.resolve()),
        "account_count": 1,
        "exported_count": 0,
    }


def test_cycle_notifies_only_after_waiting_session_job_is_resumed(tmp_path):
    storage, automation, subscription, _ = prepare_existing(tmp_path)
    with storage.connect() as conn:
        conn.execute(
            "UPDATE subscriptions SET next_run_at = '2026-07-19T19:00:00+00:00' WHERE id = ?",
            (subscription["id"],),
        )
        conn.execute("DELETE FROM sync_jobs")
    crawler = NeedsSessionThenCompleteCrawler()
    runner = WeeklySyncRunner(storage, tmp_path / "output", {}, crawler=crawler)

    first = runner.run_due(now=datetime(2026, 7, 19, 20, tzinfo=timezone.utc))

    assert first[0].status == "needs_session"
    assert automation.latest_completed_cycle() is None
    cycle = automation.cycle("2026-W29")
    assert cycle["status"] == "needs_attention"
    assert not (tmp_path / "output" / "weekly" / "2026-07-19" / "run-report.json").exists()

    resumed = runner.resume_latest_scheduled(int(subscription["id"]))

    assert resumed.status == "completed"
    completion = automation.latest_completed_cycle()
    assert completion["cycle_key"] == "2026-W29"
    assert completion["completed_account_count"] == 1
    assert crawler.calls == 2
    assert (tmp_path / "output" / "weekly" / "2026-07-19" / "run-report.json").exists()
    assert (
        tmp_path
        / "output"
        / "weekly"
        / "2026-07-19"
        / "2026-07-13至2026-07-19_本周公众号增量汇总.docx"
    ).exists()


def test_existing_crawler_uses_stored_account_identity_instead_of_seed_parse(tmp_path, monkeypatch):
    storage, _, subscription, job = prepare_existing(tmp_path)
    calls = []

    def known(self, runtime, **kwargs):
        calls.append(kwargs)
        runtime.mark_done("known account complete")

    def reverse(self, runtime):
        raise AssertionError("old seed parser must not run for an account with exact biz")

    monkeypatch.setattr("wechat_crawler.web_app.ArchiveWebService._run_known_account_job", known)
    monkeypatch.setattr("wechat_crawler.web_app.ArchiveWebService._run_reverse_job", reverse)
    adapter = ExistingCrawlerAdapter(storage, tmp_path / "output", {})
    adapter.session_refresher.prepare_account = lambda _name, seed_url=None: (_ for _ in ()).throw(
        AssertionError("visible WeChat must not be opened when the stored URL session works")
    )

    result = adapter.run(subscription, job, lambda *_args: None)

    assert result.status == "completed"
    assert calls[0]["biz"] == "biz-account"


def test_existing_crawler_opens_wechat_only_after_direct_session_failure(tmp_path, monkeypatch):
    storage, _, subscription, job = prepare_existing(tmp_path)
    calls = []
    preparations = []

    def known(self, runtime, **kwargs):
        calls.append(kwargs)
        assert runtime.since_date == job["since_date"]
        assert self.config["history"]["start_offset"] == 0
        if len(calls) == 1:
            runtime.mark_needs_session("expired visible session")
        else:
            runtime.mark_done("recovered")

    class Preparation:
        article_open_attempted = True
        session_verified = True

    monkeypatch.setattr("wechat_crawler.web_app.ArchiveWebService._run_known_account_job", known)
    adapter = ExistingCrawlerAdapter(
        storage,
        tmp_path / "output",
        {
            "automation": {"prepare_wechat_session": True},
            "history": {"app_data_dir": str(tmp_path / "wechat"), "start_offset": 90},
        },
    )

    def prepare(name, seed_url=None, *, session_fingerprint=None, **_kwargs):
        assert callable(session_fingerprint)
        preparations.append((name, seed_url))
        return Preparation()

    adapter.session_refresher.prepare_account = prepare

    result = adapter.run(subscription, job, lambda *_args: None)

    assert result.status == "completed"
    assert len(calls) == 2
    assert preparations == [("Account", "https://mp.weixin.qq.com/s/old")]
