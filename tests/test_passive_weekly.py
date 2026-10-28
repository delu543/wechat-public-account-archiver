from datetime import datetime, timedelta
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document
import pytest

from wechat_crawler.automation_store import AutomationStore
from wechat_crawler.crawler.models import FetchResult, ParsedArticle
from wechat_crawler.crawler.storage import Storage
from wechat_crawler import main as cli
from wechat_crawler.passive_weekly import PassiveWeeklyRunner
from wechat_crawler.passive_weekly_plan import make_plan, publication_day
from wechat_crawler.passive_weekly_sources import mobile_search_candidates
from wechat_crawler.passive_weekly_store import PassiveWeeklyStore

SN = "0123456789abcdef0123456789abcdef"


def url(biz="biz-account", mid=1):
    return f"https://mp.weixin.qq.com/s?__biz={biz}&mid={mid}&idx=1&sn={SN}"


def stored(storage, biz="biz-account", name="示例公众号", day="2026-09-23"):
    article = ParsedArticle(url(biz), url(biz), "旧文", name, biz, "gh_target", None, day,
                            None, "<p>历史正文</p>", "历史正文", content_hash="old-" + biz)
    ident = storage.upsert_article_url(article.source_url, "test")
    return storage.save_article(article, ident, None)


def prepare(tmp_path, *, paused=False, accounts=1):
    storage = Storage(tmp_path / "archive.sqlite")
    automation = AutomationStore(storage)
    automation.init_db()
    old = stored(storage)
    if accounts == 2:
        stored(storage, "other-account", "另一个公众号")
    automation.adopt_existing_accounts(now=datetime.fromisoformat("2026-09-22T12:00:00+08:00"))
    automation.set_automation_paused(paused)
    settings = {"delivered_baseline_max_article_id": old + accounts - 1}
    return storage, automation, settings


class Source:
    def __init__(self, *, complete=True, blocked=False, stale=False):
        self.calls, self.complete, self.blocked, self.stale = 0, complete, blocked, stale

    def read(self, accounts):
        self.calls += 1
        if self.blocked:
            return {"status": "blocked", "reason": "biz_setup_required"}
        records, notes = [], []
        for account in accounts:
            records += [{"biz": account["biz"], "url": url(account["biz"]), "publish_time": "2026-09-23"},
                        {"biz": account["biz"], "url": url(account["biz"], 2), "publish_time": "2026-10-10"}]
            notes.append({"biz": account["biz"], "schema_supported": True,
                          "all_publications_verified": self.complete,
                          "publication_proof": {"source": "native_history_index", "since": "2000-01-01", "until": "2100-01-01"},
                          "currently_followed": True if self.complete else "unknown"})
        return {"status": "read", "records": records, "accounts": notes, "coverage_measured": True,
                "source_modified_epoch": datetime.now().timestamp() - (7 * 86400 if self.stale else 0)}


class Fetcher:
    def __init__(self, *, status="ok", reason=None, wrong_biz=False, on_fetch=None):
        self.calls, self.status, self.reason, self.wrong_biz, self.on_fetch = [], status, reason, wrong_biz, on_fetch

    def fetch(self, link):
        self.calls.append(link)
        if self.on_fetch:
            self.on_fetch()
        from wechat_crawler.biz_push_reader import public_identity
        biz = "wrong-account" if self.wrong_biz else public_identity(link)[0]
        text = f'''<html><h1 id="activity-name">新文章</h1><a id="js_name">示例公众号</a>
          <em id="publish_time">2026-10-10 12:00:00</em><div id="js_content"><p>新的完整正文 {biz}</p></div>
          <script>var biz = "{biz}"; var user_name = "gh_target";</script></html>'''
        return FetchResult(link, link, 200, text, self.status, self.reason)


def runner(tmp_path, storage, settings, **kwargs):
    return PassiveWeeklyRunner(storage.db_path, tmp_path / "output", settings, **kwargs)


def test_beijing_actual_latest_not_stale_job_time_and_gap_override(tmp_path):
    storage, _, settings = prepare(tmp_path)
    with storage.connect() as conn:
        conn.execute("UPDATE articles SET publish_time='2026-10-07T21:05:14-07:00'")
    plan = make_plan(storage.db_path, settings, "2026-10-26")
    assert plan["accounts"][0]["latest_archived_beijing"] == "2026-10-08"
    assert plan["accounts"][0]["since"] == "2026-10-01"
    settings["recovery_gaps"] = {"biz-account": {"since": "2026-09-23", "until": "2026-09-28"}}
    assert make_plan(storage.db_path, settings, "2026-10-26")["accounts"][0]["since"] == "2026-09-23"
    assert publication_day("2026-10-07T21:05:14-07:00") == "2026-10-08"


def test_pause_stops_before_helper_network_tables_or_word(tmp_path):
    storage, _, settings = prepare(tmp_path, paused=True)
    source, fetcher = Source(), Fetcher()
    result = runner(tmp_path, storage, settings, sources=source, fetcher=fetcher).run("2026-10-26")
    assert result["status"] == "paused" and source.calls == 0 and not fetcher.calls
    assert not (tmp_path / "output").exists()
    with storage.connect() as conn:
        assert not conn.execute("SELECT name FROM sqlite_master WHERE name='passive_weekly_checkpoints'").fetchone()


def test_cli_plan_is_read_only_and_does_not_call_storage_init(tmp_path, monkeypatch, capsys):
    storage, _, settings = prepare(tmp_path, paused=True)
    monkeypatch.setattr(cli, "load_config", lambda path: {"database": {"path": str(storage.db_path)}, "output": {"path": str(tmp_path / "output")}, "passive_weekly": settings})
    monkeypatch.setattr(Storage, "init_db", lambda self: pytest.fail("read-only CLI must not initialize product tables"))
    assert cli.main(["weekly-run", "--plan-only", "--until", "2026-10-26"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["automation_paused"] and result["product_writes"] == 0


def test_missing_key_setup_is_one_top_level_item_and_no_writes(tmp_path):
    storage, automation, settings = prepare(tmp_path, accounts=2)
    source, fetcher = Source(blocked=True), Fetcher()
    before = [r["last_success_at"] for r in automation.list_subscriptions()]
    result = runner(tmp_path, storage, settings, sources=source, fetcher=fetcher).run("2026-10-26")
    assert result["status"] == "blocked_source" and result["required_actions"] == ["biz_setup_required"]
    assert not fetcher.calls and not (tmp_path / "output").exists()
    assert before == [r["last_success_at"] for r in automation.list_subscriptions()]
    health = runner(tmp_path, storage, settings, sources=source).inspect("2026-10-26", source_check=True)
    assert not health["coverage_measured"]
    for account in health["accounts"]:
        assert account["coverage"]["known_sample_matched"] is None
        assert account["coverage"]["known_sample_missing"] is None
        assert account["coverage"]["source_fresh"] is None


@pytest.mark.parametrize("source", [Source(complete=False), Source(stale=True)])
def test_unknown_publication_or_stale_db_cannot_be_zero_new_success(tmp_path, source):
    storage, automation, settings = prepare(tmp_path)
    fetcher = Fetcher()
    before = automation.list_subscriptions()[0]["last_success_at"]
    result = runner(tmp_path, storage, settings, sources=source, fetcher=fetcher).run("2026-10-26")
    assert result["status"] == "needs_attention" and not fetcher.calls
    assert result["accounts"][0]["status"] == "coverage_incomplete"
    assert automation.list_subscriptions()[0]["last_success_at"] == before
    assert not list((tmp_path / "output").rglob("*.docx"))


def test_complete_increment_exports_one_word_and_summary_and_repeats_safely(tmp_path):
    storage, automation, settings = prepare(tmp_path)
    source, fetcher = Source(), Fetcher()
    task = runner(tmp_path, storage, settings, sources=source, fetcher=fetcher)
    result = task.run("2026-10-26")
    assert result["status"] == "completed"
    account = result["accounts"][0]
    assert account["exported_count"] == 1 and len(fetcher.calls) == 1
    assert "2026-10-26" in Path(account["word_path"]).name
    assert [p.text for p in Document(account["word_path"]).paragraphs if p.style.name == "Heading 2"] == ["新文章"]
    assert Path(result["summary_path"]).exists()
    with storage.connect() as conn:
        assert conn.execute("SELECT verified_through FROM passive_weekly_checkpoints").fetchone()[0] == "2026-10-26"
        assert conn.execute("SELECT count(*) FROM export_batch_articles").fetchone()[0] == 2
    repeated = task.run("2026-10-26")
    assert len(fetcher.calls) == 1
    assert repeated["accounts"][0]["word_path"] == account["word_path"]


def test_explicit_partial_output_never_advances_node(tmp_path):
    storage, automation, settings = prepare(tmp_path)
    before = automation.list_subscriptions()[0]["last_success_at"]
    result = runner(tmp_path, storage, settings, sources=Source(complete=False), fetcher=Fetcher()).run("2026-10-26", allow_partial=True)
    assert result["status"] == "needs_attention" and result["accounts"][0]["word_path"]
    assert automation.list_subscriptions()[0]["last_success_at"] == before
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM passive_weekly_checkpoints").fetchone()[0] == 0


def test_word_commit_interruption_resumes_reserved_members(tmp_path, monkeypatch):
    storage, _, settings = prepare(tmp_path)
    task = runner(tmp_path, storage, settings, sources=Source(), fetcher=Fetcher())
    original = PassiveWeeklyStore.finish
    monkeypatch.setattr(PassiveWeeklyStore, "finish", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("simulated interruption")))
    first = task.run("2026-10-26")
    assert first["accounts"][0]["status"] == "failed_resumable"
    monkeypatch.setattr(PassiveWeeklyStore, "finish", original)
    second = task.run("2026-10-26")
    assert second["status"] == "completed"
    assert second["accounts"][0]["word_path"] == first["accounts"][0]["word_path"]
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM export_batches WHERE kind='weekly_incremental'").fetchone()[0] == 1


def test_shared_captcha_stops_other_accounts_and_no_success_node(tmp_path):
    storage, _, settings = prepare(tmp_path, accounts=2)
    fetcher = Fetcher(status="restricted", reason="captcha")
    result = runner(tmp_path, storage, settings, sources=Source(), fetcher=fetcher).run("2026-10-26")
    assert len(fetcher.calls) == 1
    assert [a["status"] for a in result["accounts"]] == ["blocked_shared_restriction", "blocked_shared_restriction"]
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM passive_weekly_checkpoints").fetchone()[0] == 0


def test_pause_during_fetch_prevents_article_save_and_checkpoint(tmp_path):
    storage, automation, settings = prepare(tmp_path)
    fetcher = Fetcher(on_fetch=lambda: automation.set_automation_paused(True))
    result = runner(tmp_path, storage, settings, sources=Source(), fetcher=fetcher).run("2026-10-26")
    assert result["accounts"][0]["status"] == "paused"
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM articles").fetchone()[0] == 1


def test_page_identity_independent_from_input_url_and_account_isolation(tmp_path):
    storage, _, settings = prepare(tmp_path)
    result = runner(tmp_path, storage, settings, sources=Source(), fetcher=Fetcher(wrong_biz=True)).run("2026-10-26")
    assert result["accounts"][0]["failed"] == 1
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM articles").fetchone()[0] == 1


def test_mobile_search_only_exact_biz_permanent_candidates_and_single_request():
    calls = []
    def fetch(link):
        calls.append(link)
        return FetchResult(link, link, 200, f'<a href="{url()}">x</a><a href="{url("wrong")}">y</a>', "ok")
    result = mobile_search_candidates({"biz": "biz-account", "account_name": "公开公众号"}, SimpleNamespace(fetch=fetch))
    assert len(calls) == 1 and "weixinwap?" in calls[0]
    assert len(result["records"]) == 1 and not result["complete"]


def test_no_interval_proof_cannot_advance_even_if_boolean_claims_complete(tmp_path):
    storage, _, settings = prepare(tmp_path)
    source = Source()
    original = source.read
    def read(accounts):
        report = original(accounts)
        for note in report["accounts"]:
            note.pop("publication_proof")
        return report
    source.read = read
    fetcher = Fetcher()
    result = runner(tmp_path, storage, settings, sources=source, fetcher=fetcher).run("2026-10-26")
    assert result["accounts"][0]["status"] == "coverage_incomplete" and not fetcher.calls


def test_actual_body_date_uses_beijing_and_must_match_discovery_hint(tmp_path):
    storage, _, settings = prepare(tmp_path)
    fetcher = Fetcher()
    original = fetcher.fetch
    def fetch(link):
        result = original(link)
        return FetchResult(link, link, 200, result.text.replace("2026-10-10 12:00:00", "2026-10-26T23:05:00-07:00"), "ok")
    fetcher.fetch = fetch
    result = runner(tmp_path, storage, settings, sources=Source(), fetcher=fetcher).run("2026-10-26")
    assert result["accounts"][0]["failed"] == 1
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM articles").fetchone()[0] == 1


def _three_records_source():
    source, original = Source(), Source().read
    def read(accounts):
        report = original(accounts)
        report["records"].append({"biz": "biz-account", "url": url(mid=3), "publish_time": "2026-10-10"})
        return report
    source.read = read
    return source


def _distinct_fetcher(*, fail_second=False):
    fetcher = Fetcher()
    original = fetcher.fetch
    def fetch(link):
        value = original(link)
        number = "2" if "mid=2&" in link else "3"
        if fail_second and number == "2":
            return FetchResult(link, link, None, "", "failed", "synthetic failure")
        return FetchResult(link, link, 200, value.text.replace("新文章", "新文章"+number).replace("完整正文", "完整正文"+number), "ok")
    fetcher.fetch = fetch
    return fetcher


def test_body_failure_does_not_implicitly_export_partial(tmp_path):
    storage, _, settings = prepare(tmp_path)
    result = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher(fail_second=True)).run("2026-10-26")
    assert result["accounts"][0]["failed"] == 1
    assert not result["accounts"][0]["word_path"]
    assert not list((tmp_path / "output").rglob("*.docx"))


def test_partial_retry_exports_only_the_recovered_article_in_new_batch(tmp_path):
    storage, automation, settings = prepare(tmp_path)
    first = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher(fail_second=True)).run("2026-10-26", allow_partial=True)
    assert first["accounts"][0]["exported_count"] == 1
    assert "已获取部分" in first["accounts"][0]["word_path"]
    second = runner(tmp_path, storage, settings, sources=_three_records_source(), fetcher=_distinct_fetcher()).run("2026-10-26")
    assert second["status"] == "completed" and second["accounts"][0]["exported_count"] == 1
    assert second["accounts"][0]["word_path"] != first["accounts"][0]["word_path"]
    assert not automation.unexported_articles(second["accounts"][0].get("account_id", 1))
    with storage.connect() as conn:
        assert conn.execute("SELECT count(*) FROM export_batches WHERE kind='weekly_incremental'").fetchone()[0] == 2


def test_known_stored_fragment_is_refetched_and_not_excluded_by_old_baseline(tmp_path):
    # Match the real legacy manual preview: stored before this run, but never
    # assigned completed product Word membership. Do not mutate a baseline.
    storage = Storage(tmp_path / "archive.sqlite")
    automation = AutomationStore(storage)
    automation.init_db()
    old = stored(storage)
    automation.ensure_subscription(1, url(), adopt_existing=False)
    settings = {"delivered_baseline_max_article_id": old, "incomplete_article_ids": [old]}
    fetcher, original = Fetcher(), Fetcher().fetch
    def fetch(link):
        value = original(link)
        if "mid=1&" in link:
            return FetchResult(link, link, 200, value.text.replace("2026-10-10 12:00:00", "2026-09-23 12:00:00").replace("新文章", "旧文补全"), "ok")
        return value
    fetcher.fetch = fetch
    result = runner(tmp_path, storage, settings, sources=Source(), fetcher=fetcher).run("2026-10-26")
    # The two synthetic URLs deliberately contain identical full body text.
    # Preserve account/body-hash deduplication instead of expecting two copies.
    assert result["status"] == "completed" and result["accounts"][0]["exported_count"] == 1
    document = Document(result["accounts"][0]["word_path"])
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 2"] == ["旧文补全"]
